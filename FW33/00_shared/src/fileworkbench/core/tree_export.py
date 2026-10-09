"""ディレクトリのスナップショットとTXT/CSV/Excel/HTML出力。"""
import csv
import html
import os
import tempfile
from pathlib import Path
from datetime import datetime

COLUMNS=[('number','番号'),('depth','階層'),('relative_path','相対パス'),('kind','種類'),('size','サイズ(bytes)'),('modified','更新日時'),('status','状態'),('error','エラー')]

def scan_tree(root,cancel=lambda:False):
    root=Path(root).absolute()
    if not root.is_dir():raise NotADirectoryError(str(root))
    rows=[]
    def row(path,kind,stat=None,error=''):
        relative=path.relative_to(root)
        return dict(path=str(path),name=path.name or str(path),relative_path='.' if path==root else str(relative),depth=0 if path==root else len(relative.parts),kind=kind,size=stat.st_size if stat and kind=='ファイル' else '',modified=datetime.fromtimestamp(stat.st_mtime).isoformat(timespec='seconds') if stat else '',status='失敗' if error else '成功',error=error)
    rows.append(row(root,'フォルダー',root.stat()));pending=[root]
    while pending:
        if cancel():break
        directory=pending.pop()
        try:
            with os.scandir(directory) as iterator:entries=sorted(iterator,key=lambda e:(not e.is_dir(follow_symlinks=False),e.name.casefold()))
        except OSError as exc:
            for record in rows:
                if record['path']==str(directory):record['status']='失敗';record['error']=str(exc);break
            continue
        for entry in entries:
            path=Path(entry.path)
            try:
                stat=entry.stat(follow_symlinks=False);link=entry.is_symlink() or getattr(stat,'st_reparse_tag',0) in (0xA0000003,0xA000000C);folder=entry.is_dir(follow_symlinks=False)
                kind='リンク' if link else 'フォルダー' if folder else 'ファイル'
                rows.append(row(path,kind,stat))
                if folder and not link:pending.append(path)
            except OSError as exc:rows.append(row(path,'不明',error=str(exc)))
    if cancel():rows[0]['status']='停止';rows[0]['error']='利用者が停止したため途中結果です'
    rows.sort(key=lambda r:(r['depth'],r['relative_path'].casefold()))
    return rows

def export_tree(rows,root,destination,format_name,include_files=True,numbering=False,overwrite=False):
    """一時ファイルへ書き、成功後に保存先を置き換える。"""
    root=Path(root);destination=Path(destination)
    chosen=[dict(r) for r in rows if include_files or r['kind']=='フォルダー']
    for i,r in enumerate(chosen,1):r['number']=i if numbering else ''
    handle=tempfile.NamedTemporaryFile(prefix='.fwn_export_',suffix=destination.suffix,dir=destination.parent,delete=False);temporary=Path(handle.name);handle.close()
    try:
        if format_name=='CSV':
            with temporary.open('w',newline='',encoding='utf-8-sig') as out:
                w=csv.writer(out);w.writerow([title for _,title in COLUMNS]);w.writerows([[r[k] for k,_ in COLUMNS] for r in chosen])
        elif format_name=='Excel':
            from openpyxl import Workbook
            from openpyxl.styles import Font,PatternFill,Alignment
            wb=Workbook();ws=wb.active;ws.title='構成一覧';ws.append([title for _,title in COLUMNS])
            for r in chosen:
                ws.append([r[k] for k,_ in COLUMNS])
                for cell in ws[ws.max_row]:
                    # ファイル名等が数式として評価されないよう文字列型を固定。
                    if isinstance(cell.value,str):cell.data_type='s'
            for cell in ws[1]:cell.font=Font(bold=True,color='205FA0');cell.fill=PatternFill('solid',fgColor='DCEEFF')
            widths=[10,8,65,18,20,24,12,70]
            for i,width in enumerate(widths,1):ws.column_dimensions[chr(64+i)].width=width
            ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;wb.save(temporary);wb.close()
        elif format_name in ('HTML','TXT'):
            children={}
            for r in chosen:
                if r['relative_path']=='.':continue
                parent=str(Path(r['relative_path']).parent)
                children.setdefault(parent,[]).append(r)
            for entries in children.values():entries.sort(key=lambda r:(r['kind']!='フォルダー',r['name'].casefold()))
            def title(r):return (str(r['number'])+'. ' if numbering else '')+r['name']
            output=[]
            if format_name=='TXT':
                output.append(('1. ' if numbering else '')+str(root))
                stack=[(r,'',i==len(children.get('.',[]))-1) for i,r in reversed(list(enumerate(children.get('.',[]))))]
                while stack:
                    r,prefix,last=stack.pop();output.append(prefix+('└─ ' if last else '├─ ')+title(r)+(' [失敗: '+r['error']+']' if r['error'] else ''))
                    nested=children.get(r['relative_path'],[]);newprefix=prefix+('   ' if last else '│  ')
                    stack.extend((child,newprefix,i==len(nested)-1) for i,child in reversed(list(enumerate(nested))))
                temporary.write_text('\n'.join(output)+'\n',encoding='utf-8-sig')
            else:
                escape=html.escape;output=['<!doctype html><html lang="ja"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>フォルダー構成</title><style>body{font-family:system-ui,sans-serif;background:#f7faff;color:#31445f;margin:24px}details{margin:6px 0}ul{list-style:none;padding-left:24px}li{padding:4px}a{color:#205fa0;text-decoration:none}.error{color:#a14d66}summary{cursor:pointer}</style>',f'<h1>フォルダー構成</h1><p>{escape(str(root))}</p><p>フォルダー名をクリックして折り畳み、リンクから元ファイルを開けます。リンクは元の保存場所を参照します。</p><ul>']
                # 閉じタグもスタックへ積み、深いツリーでも再帰制限に依存しない。
                stack=list(reversed(children.get('.',[])))
                while stack:
                    r=stack.pop()
                    if isinstance(r,str):output.append(r);continue
                    link=Path(r['path']).absolute().as_uri();label=escape(title(r));error=f' <span class="error">{escape(r["error"])}</span>' if r['error'] else ''
                    if r['kind']=='フォルダー':
                        output.append(f'<li><details open><summary>📁 {label}{error}</summary><a href="{escape(link,quote=True)}">フォルダーを開く</a><ul>');stack.append('</ul></details></li>');stack.extend(reversed(children.get(r['relative_path'],[])))
                    else:output.append(f'<li>📄 <a href="{escape(link,quote=True)}">{label}</a>{error}</li>')
                output.append('</ul></html>');temporary.write_text('\n'.join(output),encoding='utf-8')
        else:raise ValueError('未対応の出力形式: '+format_name)
        if overwrite:os.replace(temporary,destination)
        else:
            from .safety import no_replace
            no_replace(temporary,destination)
    finally:temporary.unlink(missing_ok=True)
    return str(destination)
