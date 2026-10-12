"""Safe content readers that other apps can use without importing a preview UI."""
from pathlib import Path
import codecs
import csv,io,zipfile,xml.etree.ElementTree as ET

MEDIA={'.mp3','.wav','.flac','.ogg','.aac','.m4a','.wma','.mp4','.mkv','.avi','.mov','.webm','.wmv','.m4v','.mpeg','.mpg','.opus'}
TEXT={'.txt','.md','.log','.csv','.tsv','.json','.xml','.yaml','.yml','.ini','.cfg','.py','.cs','.js','.ts','.html','.css','.bat','.ps1','.sql','.svg'}

class PreviewReaderPort:
    """Readers expose supports(path) and read(path) -> text."""
    def supports(self,path):raise NotImplementedError
    def read(self,path):raise NotImplementedError

class TextReader(PreviewReaderPort):
    def supports(self,path):return path.suffix.lower() in TEXT
    def read(self,path):
        with path.open('rb') as f:raw=f.read(1024*1024+1)
        truncated=len(raw)>1024*1024;raw=raw[:1024*1024]
        if raw.startswith((codecs.BOM_UTF16_LE,codecs.BOM_UTF16_BE)):
            text=raw.decode('utf-16')
        else:
            for encoding in ('utf-8-sig','cp932'):
                try:text=raw.decode(encoding);break
                except UnicodeError:continue
            else:text=raw.decode('utf-8',errors='replace')
        return text+('\n\n[先頭1MiBのみ表示]' if truncated else '')

class OfficeReader(PreviewReaderPort):
    def supports(self,path):return path.suffix.lower() in ('.docx','.pptx','.xlsx')
    def read(self,path):
        if path.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True)
            try:
                lines=['[表の値を簡易表示。書式・図形は対象外]']
                for sheet in book.worksheets[:10]:
                    lines.append('\n## '+sheet.title)
                    for row in sheet.iter_rows(max_row=200,max_col=20,values_only=True):lines.append('\t'.join('' if v is None else str(v) for v in row))
                return '\n'.join(lines)
            finally:book.close()
        with zipfile.ZipFile(path) as z:
            names=['word/document.xml'] if path.suffix.lower()=='.docx' else sorted(n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
            lines=['[本文を簡易表示。原稿のページレイアウト・図形は対象外]']
            total=0
            for name in names[:200]:
                size=z.getinfo(name).file_size;total+=size
                if total>10*1024*1024:lines.append('[読込上限]');break
                xml=ET.fromstring(z.read(name));lines.append('\n'+name+'\n'+'\n'.join(e.text for e in xml.iter() if e.tag.endswith('}t') and e.text))
            return '\n'.join(lines)

class ArchiveReader(PreviewReaderPort):
    def supports(self,path):return path.suffix.lower()=='.zip'
    def read(self,path):
        with zipfile.ZipFile(path) as z:return '[ZIP内一覧。展開・実行はしません]\n'+'\n'.join(f'{i.filename}\t{i.file_size:,} bytes' for i in z.infolist()[:3000])

class MetadataReader(PreviewReaderPort):
    def supports(self,path):return True
    def read(self,path):return f'専用プレビューのない形式です。\n\n名前: {path.name}\n種類: {path.suffix}\nサイズ: {path.stat().st_size:,} bytes\n場所: {path}\n\n右クリックから既定アプリで開けます。'

READERS=[OfficeReader(),ArchiveReader(),TextReader(),MetadataReader()]
def read_content(path):
    path=Path(path)
    return next(reader for reader in READERS if reader.supports(path)).read(path)
